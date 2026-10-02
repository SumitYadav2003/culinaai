"""
Export ingredient prices from the price workbook into the CSV the app loads
with `python manage.py load_prices`.

Run from the project root (development only):
    pip install openpyxl
    python scripts/export_prices.py [docs/evaluation/CulinaAI_ingredient_prices.xlsx]

The workbook must have been saved by Excel or LibreOffice after its last edit,
so that the "Price per kg" formula cells hold their values.

What goes in the CSV, per food code:
    price_per_kg_gbp  the price per kg, or empty when there is no price yet
    source            where it came from (ONS / shop average / tap water)
    status            the workbook's status column, e.g. "ONS stand-in, check"
    note              the workbook's note (ranges, weights used, why blank)

Every shop price behind a "Shop average" is in the workbook's "Shop prices" tab,
which is also exported, so each average can be checked.
"""

import csv
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = PROJECT_ROOT / "docs" / "evaluation" / "CulinaAI_ingredient_prices.xlsx"
DATA_DIR = PROJECT_ROOT / "backend" / "nutrition" / "data"
PRICES_CSV = DATA_DIR / "ingredient_prices.csv"
EVIDENCE_CSV = DATA_DIR / "ingredient_prices_shop_evidence.csv"

SOURCE_FOR_STATUS = {
    "ONS, same item": "ONS average price, Aug 2026",
    "ONS stand-in, check": "ONS average price, Aug 2026 (nearest item)",
    "ONS + egg weight": "ONS average price, Aug 2026",
    "Shop average": "Average of {shops}, 2 Oct 2026",
    "Shop stand-in, check": "Average of {shops}, 2 Oct 2026 (nearest item)",
    "Free": "Tap water",
    "No price yet": "",
}


def export_prices(sheet):
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        code, name, _, status = values[0], values[1], values[2], values[3]
        if not code or code == "(example)":
            continue
        if status not in SOURCE_FOR_STATUS:
            raise SystemExit(f"Unknown status {status!r} for {code}")
        per_kg = values[8]
        if status == "No price yet":
            per_kg = ""
        elif not isinstance(per_kg, (int, float)):
            raise SystemExit(f"{code} has status {status!r} but no price per kg. Open and re-save the workbook.")
        rows.append({
            "food_code": code,
            "food_name": name,
            "price_per_kg_gbp": "" if per_kg == "" else round(per_kg, 3),
            "source": SOURCE_FOR_STATUS[status].format(shops=values[9] or ""),
            "status": status,
            "note": values[11] or "",
        })
    return rows


def export_evidence(sheet):
    header = [cell.value for cell in sheet[1]]
    return header, [list(values) for values in sheet.iter_rows(min_row=2, values_only=True) if values[0]]


def main():
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)

    prices = export_prices(workbook["Prices"])
    with PRICES_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(prices[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(prices)

    header, evidence = export_evidence(workbook["Shop prices"])
    with EVIDENCE_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(evidence)

    priced = sum(1 for row in prices if row["price_per_kg_gbp"] != "")
    print(f"Prices: {len(prices)} foods, {priced} with a price -> {PRICES_CSV}")
    print(f"Shop evidence: {len(evidence)} rows -> {EVIDENCE_CSV}")


if __name__ == "__main__":
    main()
