"""
Extract the nutrients CulinaAI needs from the official CoFID 2021 Excel file
into a small, clean CSV that the app loads with `python manage.py load_cofid`.

Run from the project root (development only):
    pip install openpyxl
    python scripts/extract_cofid.py path/to/McCance_Widdowsons_..._2021.xlsx

Source: McCance and Widdowson's Composition of Foods Integrated Dataset 2021,
Public Health England, Open Government Licence v3.0.
https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid

How CoFID values are cleaned:
    a number  -> kept as it is
    "Tr"      -> 0     (trace: present only in a tiny amount)
    "N"       -> empty (CoFID has no reliable figure; never treated as zero)
    blank     -> empty
Salt is not in CoFID, so it is worked out from sodium: salt (g) = sodium (mg) x 2.5 / 1000.
Fibre uses AOAC fibre (the method on UK food labels). If a food has no AOAC
value, the older NSP fibre value is used instead and the row is counted in the report.
"""

import csv
import re
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_CSV = PROJECT_ROOT / "backend" / "nutrition" / "data" / "cofid_2021.csv"

# Output column -> CoFID header names to look for (lowercase, spaces collapsed).
WANTED_COLUMNS = {
    "energy_kcal": ["energy (kcal) (kcal)", "energy (kcal)"],
    "energy_kj": ["energy (kj) (kj)", "energy (kj)"],
    "protein_g": ["protein (g)"],
    "fat_g": ["fat (g)"],
    "saturates_g": ["satd fa /100g fd (g)", "saturated fatty acids per 100g food (g)"],
    "carbohydrate_g": ["carbohydrate (g)"],
    "sugars_g": ["total sugars (g)"],
    "aoac_fibre_g": ["aoac fibre (g)"],
    "nsp_fibre_g": ["nsp (g)"],
    "sodium_mg": ["sodium (mg)"],
}

OUTPUT_FIELDS = [
    "food_code", "name", "food_group", "energy_kcal", "energy_kj", "protein_g", "fat_g",
    "saturates_g", "carbohydrate_g", "sugars_g", "fibre_g", "salt_g",
]


def clean_header(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def clean_value(value):
    """Return a float, 0.0 for trace, or None when CoFID has no figure."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text == "" or text.upper() == "N":
        return None
    if text.lower() == "tr":
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def find_header_row(sheet):
    """
    The header row is the first row with a 'Food Code' cell.
    In CoFID 2021 the '1.4 Inorganics' sheet leaves that cell blank, so a row
    with 'Food Name' and an empty first cell counts too (codes are in column A).
    """
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        headers = [clean_header(cell) for cell in row]
        if "food code" in headers:
            return row_number, headers
        if "food name" in headers and headers and headers[0] == "":
            headers[0] = "food code"
            return row_number, headers
    return None, None


def read_sheets(workbook):
    """
    Read every sheet that has a 'Food Code' column.
    Returns: foods {code: {"name", "group"}}, values {code: {output_column: value}}, found {output_column: sheet}
    """
    foods, values, found = {}, {}, {}
    for sheet in workbook.worksheets:
        header_row, headers = find_header_row(sheet)
        if header_row is None:
            continue

        code_col = headers.index("food code")
        name_col = headers.index("food name") if "food name" in headers else None
        group_col = headers.index("group") if "group" in headers else None

        columns_here = {}
        for output, options in WANTED_COLUMNS.items():
            if output in found:
                continue
            for option in options:
                if option in headers:
                    columns_here[output] = headers.index(option)
                    found[output] = f"{sheet.title} / {option}"
                    break

        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            code = str(row[code_col] or "").strip()
            if not re.match(r"^\d{2}-\d{3}$", code):  # CoFID codes look like 11-123
                continue
            if name_col is not None and code not in foods:
                foods[code] = {
                    "name": str(row[name_col] or "").strip(),
                    "group": str(row[group_col] or "").strip() if group_col is not None else "",
                }
            for output, col in columns_here.items():
                values.setdefault(code, {})[output] = clean_value(row[col])
    return foods, values, found


def build_rows(foods, values):
    rows, used_nsp = [], 0
    for code, info in sorted(foods.items()):
        food = values.get(code, {})
        fibre = food.get("aoac_fibre_g")
        if fibre is None and food.get("nsp_fibre_g") is not None:
            fibre = food["nsp_fibre_g"]
            used_nsp += 1
        sodium = food.get("sodium_mg")
        rows.append({
            "food_code": code,
            "name": info["name"],
            "food_group": info["group"],
            "energy_kcal": food.get("energy_kcal"),
            "energy_kj": food.get("energy_kj"),
            "protein_g": food.get("protein_g"),
            "fat_g": food.get("fat_g"),
            "saturates_g": food.get("saturates_g"),
            "carbohydrate_g": food.get("carbohydrate_g"),
            "sugars_g": food.get("sugars_g"),
            "fibre_g": fibre,
            "salt_g": round(sodium * 2.5 / 1000, 3) if sodium is not None else None,
        })
    return rows, used_nsp


def main():
    if len(sys.argv) != 2:
        sys.exit("Usage: python scripts/extract_cofid.py path/to/cofid_2021.xlsx")

    workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    foods, values, found = read_sheets(workbook)

    print("Columns found:")
    for output in WANTED_COLUMNS:
        print(f"  {output:15s} <- {found.get(output, 'NOT FOUND')}")
    missing = [o for o in WANTED_COLUMNS if o not in found and o != "nsp_fibre_g"]
    if missing:
        sys.exit(f"Stopping: could not find {missing}. Check the header names in the Excel file.")

    rows, used_nsp = build_rows(foods, values)
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    print(f"\nWrote {len(rows)} foods to {OUTPUT_CSV.relative_to(PROJECT_ROOT)}")
    print(f"Foods using NSP fibre because AOAC was missing: {used_nsp}")
    for field in OUTPUT_FIELDS[3:]:
        empty = sum(1 for r in rows if r[field] is None)
        print(f"  {field:15s} empty in {empty} foods")


if __name__ == "__main__":
    main()
